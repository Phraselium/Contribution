"""M1-F7 · tablas_m1.xlsx con la convención de colores del plan: azul = dato introducido (con su fuente en una nota de la celda),
negro = calculado en la hoja (fórmula), verde = viene de otra pestaña (fórmula que referencia otra hoja). Sin paneles
inmovilizados. Sin red. Salida: output/f7/tablas_m1.xlsx (publicar.py la copia a publicacion/documento_m1/).
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from comun import OUT  # noqa: E402

AZUL, NEGRO, VERDE = Font(color="1F4E9A"), Font(color="000000"), Font(color="2E7D32")
CAB = Font(bold=True)
FONDO = PatternFill("solid", fgColor="EEF2F7")
LINEA = Border(bottom=Side(style="thin", color="999999"))


def cabecera(ws, fila, textos):
    for j, t in enumerate(textos, 1):
        c = ws.cell(fila, j, t)
        c.font, c.fill, c.border = CAB, FONDO, LINEA
        c.alignment = Alignment(wrap_text=True, vertical="center")


def dato(ws, f, c, v, fuente):
    cel = ws.cell(f, c, None if pd.isna(v) else v)
    cel.font = AZUL
    cel.comment = Comment(f"Fuente: {fuente}", "M1")
    return cel


def anchos(ws, w):
    for j, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(j)].width = x


def main() -> int:
    wb = Workbook()
    # Leyenda
    ws = wb.active
    ws.title = "Leyenda"
    ws["A1"], ws["A1"].font = "tablas_m1.xlsx · Estudio M1 (solo información pública)", Font(bold=True, size=12)
    for i, (t, f) in enumerate([("Dato introducido (fuente en la nota de la celda)", AZUL), ("Calculado en la hoja", NEGRO),
                                ("Viene de otra pestaña", VERDE)], 3):
        ws.cell(i, 1, t).font = f
    ws["A7"] = "Todas las cifras difundibles proceden de output/cifras_clave.csv (pestaña «Cifras»)."
    anchos(ws, [80])
    # Fiscal: componentes por persona (dato) y saldos calculados
    r = pd.read_csv(OUT / "f4" / "neta_por_grupo_2024.csv").set_index("grupo")
    ws = wb.create_sheet("Fiscal_2024")
    comps = [c for c in r.columns if c.startswith(("ing_", "gas_"))]
    cabecera(ws, 1, ["Partida (€ por persona y año, 2024)"] + list(r.index))
    fuente = "output/f4/neta_por_grupo_2024.csv (ECV 2025, EPF, ENSE y cuentas de las AAPP de Eurostat; escenario principal)"
    for i, c in enumerate(comps, 2):
        ws.cell(i, 1, c)
        for j, g in enumerate(r.index, 2):
            dato(ws, i, j, round(float(r.loc[g, c]), 1), fuente)
    n = len(comps)
    fi = [i for i, c in enumerate(comps, 2) if c.startswith("ing_")]
    fg = [i for i, c in enumerate(comps, 2) if c.startswith("gas_")]
    for k, (etq, filas) in enumerate((("Ingresos", fi), ("Gastos", fg)), n + 2):
        ws.cell(k, 1, etq).font = CAB
        for j in range(2, 2 + len(r.index)):
            L = get_column_letter(j)
            ws.cell(k, j, f"=SUM({L}{min(filas)}:{L}{max(filas)})").font = NEGRO
    ws.cell(n + 4, 1, "Contribución neta").font = CAB
    for j in range(2, 2 + len(r.index)):
        L = get_column_letter(j)
        ws.cell(n + 4, j, f"={L}{n + 2}-{L}{n + 3}").font = NEGRO
    anchos(ws, [38, 16, 16, 18])
    # Resumen que toma la neta de la hoja Fiscal (verde)
    ws2 = wb.create_sheet("Resumen")
    cabecera(ws2, 1, ["Grupo", "Neta por persona (€)", "Población", "Neta total (M€)"])
    for i, g in enumerate(r.index, 2):
        L = get_column_letter(i)
        ws2.cell(i, 1, g)
        ws2.cell(i, 2, f"=Fiscal_2024!{L}{n + 4}").font = VERDE
        dato(ws2, i, 3, round(float(r.loc[g, "poblacion"])), "ECV 2025 (INE), suma de pesos RB050")
        ws2.cell(i, 4, f"=B{i}*C{i}/1000000").font = NEGRO
    anchos(ws2, [20, 20, 16, 18])
    # Empleo
    e = pd.read_csv(OUT / "f3" / "empleo_nacional.csv")
    e = e[(e.origen_def == "nacimiento") & (e.anio == 2025)]
    ws3 = wb.create_sheet("Empleo_2025")
    cabecera(ws3, 1, ["Grupo (nacimiento)", "Población 16-64", "Tasa de empleo 16-64", "Tasa de paro (16 años o más)",
                      "Temporalidad (asalariados de 16 años o más)", "Ocupados de 16 años o más (no dividir por la población 16-64)",
                      "Cuota de los ocupados de 16 años o más"])
    for i, (_, x) in enumerate(e.iterrows(), 2):
        ws3.cell(i, 1, x.grupo)
        for j, v in enumerate(["pob_16_64", "tasa_empleo", "tasa_paro", "temporalidad", "ocupados"], 2):
            dato(ws3, i, j, round(float(x[v]), 4), "EPA (INE), microdatos 2025, medias anuales; output/f3/empleo_nacional.csv")
        ws3.cell(i, 7, f"=F{i}/SUM($F$2:$F${1 + len(e)})").font = NEGRO
    anchos(ws3, [26, 16, 14, 14, 14, 14, 18])
    # Cifras clave
    c = pd.read_csv(OUT / "cifras_clave.csv")
    ws4 = wb.create_sheet("Cifras")
    cols = ["id", "valor", "ic_inf", "ic_sup", "unidad", "capa", "clasificacion", "periodo", "ambito", "grupo_origen", "fuente", "descripcion"]
    cabecera(ws4, 1, cols)
    for i, (_, x) in enumerate(c.iterrows(), 2):
        for j, k in enumerate(cols, 1):
            v = x[k]
            cel = ws4.cell(i, j, None if pd.isna(v) else v)
            if k in ("valor", "ic_inf", "ic_sup"):
                cel.font = AZUL
                cel.comment = Comment(f"Fuente: {x['fuente']} ({x['script']})", "M1")
    anchos(ws4, [44, 10, 10, 10, 16, 6, 14, 12, 14, 18, 40, 80])
    (OUT / "f7").mkdir(parents=True, exist_ok=True)
    wb.save(OUT / "f7" / "tablas_m1.xlsx")
    print("f7_tablas: output/f7/tablas_m1.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
